"""Insert the motion-graphics slides (figures/anim/*.mp4) into the talk deck.

    python animations.py          # render the videos first
    python add_animations.py      # edits talk/differentiation_integration_talk.pptx in place

Each animation gets its own 'NLD Title Only' slide right after the static
slide it illustrates. Slides already carrying an animation are skipped, so the
script can be re-run safely.
"""
import os
import copy
import subprocess
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..')
DECK = os.path.join(ROOT, 'talk', 'differentiation_integration_talk.pptx')
ANIM = os.path.join(ROOT, 'figures', 'anim')

# (after the slide whose title starts with ..., eyebrow, title, caption, notes, video)
SLIDES = [
    ('Why it is ill-posed', 'THE PROBLEM',
     'Differentiation and integration: one gain, opposite ends',
     'Animated. Left: noise gain ωⁿ as the order n sweeps from two integrals to two derivatives. '
     'Right: signal and 1% position noise after n-fold differentiation; shaded where noise exceeds signal.',
     'Each derivative multiplies the spectrum by i omega, each integral divides by it. White position noise is flat; '
     'the signal falls off above the modes. Differentiating lifts the noise above the signal at high frequency; '
     'integrating lifts low-frequency errors (offsets, drift) instead. Same problem, opposite directions.',
     'illposed'),
    ('The FFT caveat', 'DIFFERENTIATION',
     'Why the FFT fails at the ends of the record',
     'Animated. The FFT treats the 3 s record as one period, so the record end meets its own start with a jump. '
     'Spectral velocity error reaches 2.7σ at the last sample; the state-space smoother stays below 0.02σ.',
     'A damped decay does not return to its starting value, so the periodic extension has a discontinuity. '
     'Its Gibbs ringing lands on the first and last few cycles. Detrending and tapering reduce it, but a method '
     'with no periodicity assumption (state-space, Whittaker) has no such problem.',
     'edges'),
    ('Every method has an optimum', 'DIFFERENTIATION',
     'One knob: sweeping the cutoff',
     'Animated. Whittaker acceleration from 1% noisy position as the cutoff fc sweeps from 20 to 600 Hz, then '
     'returns to the optimum. Right: error against exact ground truth.',
     'Below the optimum the estimate is biased: the modes are smoothed away. Above it, noise amplified as omega '
     'squared dominates. The optimum depends on the data, which is why self-tuning (GCV, maximum likelihood) matters.',
     'knob'),
    ('Vibro-impact oscillator', 'DIFFERENTIATION',
     'At an impact, the L1 penalty tracks the jump most closely',
     'Animated. Zoom onto one impact of the vibro-impact oscillator, 0.5% noise. '
     'Legend: velocity error within ±4 ms of all five impacts.',
     'The velocity genuinely jumps at contact. Quadratic penalties (Whittaker, spectral cutoff) spread the jump and ring; '
     'total variation allows a sharp change. The gain is real but modest: about 20% lower error near impacts.',
     'impact'),
    ('If you have both sensors', 'INTEGRATION',
     'Forward filter, then backward smoother',
     'Animated. One state-space model fuses 1% noisy displacement with a drifting accelerometer. '
     'Forward Kalman pass: 1.70% velocity error. Backward RTS pass: 0.43%. Grey: ±2σ posterior band.',
     'The causal filter uses only the past, so its band is wider. The RTS pass adds the future and shrinks the band '
     'everywhere except near the ends. The band comes from the fitted noise model: the estimator reports its own uncertainty.',
     'fusion'),
]


def title_of(slide):
    for sh in slide.placeholders:
        if sh.placeholder_format.type is not None and 'TITLE' in str(sh.placeholder_format.type):
            return sh.text_frame.text
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith('Title'):
            return sh.text_frame.text
    return ''


def caption_box(slide):
    for sh in slide.shapes:
        if sh.shape_type == 17 and sh.name.startswith('TextBox'):
            return sh
    return None


def main():
    prs = Presentation(DECK)
    layout = next(l for l in prs.slide_layouts if l.name == 'NLD Title Only')
    have = {sh.name for s in prs.slides for sh in s.shapes}
    template_cap = next(caption_box(s) for s in prs.slides if caption_box(s) is not None)
    ids = prs.slides._sldIdLst
    for anchor, eyebrow, title, caption, notes, name in SLIDES:
        if 'anim:' + name in have:
            print('skip', name); continue
        pos = next(i for i, s in enumerate(prs.slides) if anchor.lower() in
                   (title_of(s) + ' ' + ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)).lower())
        ref = prs.slides[pos]
        s = prs.slides.add_slide(layout)
        for ph in s.placeholders:
            nm = ph.name.lower()
            if 'title' in nm:
                ph.text_frame.text = title
            elif 'text' in nm:
                ph.text_frame.text = eyebrow
        pic = next(sh for sh in ref.shapes if sh.shape_type == 13)
        W = pic.width; H = int(W*700/1600)
        top = pic.top + (pic.height - H)//2
        poster = os.path.join(ANIM, name + '.jpg')
        if not os.path.exists(poster):
            subprocess.run(['ffmpeg', '-v', 'error', '-y', '-sseof', '-0.1', '-i', os.path.join(ANIM, name + '.mp4'),
                            '-frames:v', '1', '-q:v', '3', poster], check=True)
        mv = s.shapes.add_movie(os.path.join(ANIM, name + '.mp4'), pic.left, top, W, H,
                                poster_frame_image=poster, mime_type='video/mp4')
        mv.name = 'anim:' + name
        cap = copy.deepcopy(template_cap._element)
        s.shapes._spTree.append(cap)
        tb = s.shapes[-1]
        p0 = tb.text_frame.paragraphs[0]
        for p in tb.text_frame.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        for r in p0.runs[1:]:
            r._r.getparent().remove(r._r)
        p0.runs[0].text = caption
        ns = s.notes_slide
        if ns.notes_text_frame is None and ref.has_notes_slide:   # notes master lacks a body placeholder
            body = next(sh for sh in ref.notes_slide.placeholders if sh.placeholder_format.type == 2)
            ns.shapes._spTree.append(copy.deepcopy(body._element))
        if ns.notes_text_frame is not None:
            ns.notes_text_frame.text = notes
        el = ids[-1]
        ids.remove(el); ids.insert(pos + 1, el)
        print('added', name, 'after slide', pos + 1)
    prs.save(DECK)


if __name__ == '__main__':
    main()
