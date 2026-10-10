
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK   = RGBColor(0x1A, 0x1A, 0x1A)
ACC   = RGBColor(0x2F, 0x6F, 0x9F)
WARN  = RGBColor(0xC2, 0x45, 0x3C)
GOOD  = RGBColor(0x1B, 0x7F, 0x5C)
MUTE  = RGBColor(0x6B, 0x72, 0x80)
BG    = RGBColor(0xFF, 0xFF, 0xFF)
FONT  = "Helvetica Neue"
W, H  = 13.333, 7.5

def _tb(slide, x, y, w, h, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def _run(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = FONT
    return r

def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _rule(slide, y=1.18, x=0.62, w=12.1, color=ACC, h=0.035):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background(); s.shadow.inherit = False
    return s

def _head(slide, title, sub=None):
    tf = _tb(slide, 0.62, 0.42, 12.1, 0.8)
    p = tf.paragraphs[0]; _run(p, title, 30, bold=True)
    _rule(slide)
    if sub:
        tf2 = _tb(slide, 0.62, 1.30, 12.1, 0.45)
        _run(tf2.paragraphs[0], sub, 16, color=MUTE, italic=True)

def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def title_slide(prs, title, subtitle, author, date):
    s = _blank(prs)
    tf = _tb(s, 0.9, 2.35, 11.6, 1.6)
    _run(tf.paragraphs[0], title, 42, bold=True)
    _rule(s, y=4.05, x=0.92, w=5.2)
    tf = _tb(s, 0.92, 4.32, 11.4, 1.5)
    _run(tf.paragraphs[0], subtitle, 20, color=MUTE)
    p = tf.add_paragraph(); p.space_before = Pt(22)
    _run(p, author, 16)
    p = tf.add_paragraph(); _run(p, date, 14, color=MUTE)
    return s

def section(prs, n, title, blurb):
    s = _blank(prs)
    tf = _tb(s, 0.9, 2.7, 11.5, 0.7)
    _run(tf.paragraphs[0], n, 16, bold=True, color=ACC)
    tf = _tb(s, 0.9, 3.15, 11.5, 1.0)
    _run(tf.paragraphs[0], title, 34, bold=True)
    _rule(s, y=4.22, x=0.92, w=4.2)
    tf = _tb(s, 0.92, 4.5, 10.8, 1.0)
    _run(tf.paragraphs[0], blurb, 17, color=MUTE)
    return s

def bullets(prs, title, items, notes="", sub=None, top=1.85, size=19):
    s = _blank(prs); _head(s, title, sub)
    tf = _tb(s, 0.72, top, 11.9, 7.5 - top - 0.5)
    first = True
    for it in items:
        lvl, txt = (it if isinstance(it, tuple) else (0, it))
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(11 if lvl == 0 else 6)
        p.level = lvl
        if lvl == 0:
            _run(p, "\u2014  ", size, color=ACC, bold=True)
            _run(p, txt, size)
        else:
            p.space_before = Pt(2)
            r = _run(p, "      " + txt, size - 3, color=MUTE)
    _notes(s, notes)
    return s

def figure(prs, title, img, caption=None, notes="", full=True):
    s = _blank(prs); _head(s, title)
    from PIL import Image
    iw, ih = Image.open(img).size
    ar = iw / ih
    avail_w, avail_h = 12.1, (5.02 if caption else 5.65)
    w = avail_w; h = w / ar
    if h > avail_h:
        h = avail_h; w = h * ar
    x = (W - w) / 2; y = 1.5 + (avail_h - h) / 2
    s.shapes.add_picture(img, Inches(x), Inches(y), width=Inches(w))
    if caption:
        tf = _tb(s, 0.86, 6.66, 11.6, 0.6)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        _run(p, caption, 13, color=MUTE, italic=True)
    _notes(s, notes)
    return s

def proscons(prs, title, idea, pros, cons, verdict, notes="", vcolor=ACC):
    s = _blank(prs); _head(s, title)
    tf = _tb(s, 0.72, 1.62, 11.9, 0.8)
    _run(tf.paragraphs[0], idea, 18, color=INK, italic=True)
    nmax = max(len(pros), len(cons))
    bs = 17 if nmax <= 4 else 15
    sa = 9 if nmax <= 4 else 6
    for col, (hdr, items, colr) in enumerate([("Strengths", pros, GOOD), ("Limitations", cons, WARN)]):
        x = 0.72 + col * 6.15
        tf = _tb(s, x, 2.62, 5.75, 0.4)
        _run(tf.paragraphs[0], hdr.upper(), 14, bold=True, color=colr)
        tf = _tb(s, x, 3.08, 5.75, 3.0)
        first = True
        for it in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            p.space_after = Pt(sa)
            _run(p, "\u2022  ", bs, color=colr, bold=True)
            _run(p, it, bs)
    from pptx.enum.shapes import MSO_SHAPE
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.28), Inches(11.9), Inches(0.82))
    bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor(0xF2, 0xF5, 0xF8)
    bar.line.fill.background(); bar.shadow.inherit = False
    tf = bar.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.24); tf.margin_top = Inches(0.1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    _run(p, "USE IT WHEN   ", 13, bold=True, color=vcolor)
    _run(p, verdict, 16)
    _notes(s, notes)
    return s

def table(prs, title, header, rows, widths, notes="", sub=None, fs=14, hfs=13):
    s = _blank(prs); _head(s, title, sub)
    top = 1.95 if sub else 1.72
    nr, nc = len(rows) + 1, len(header)
    gt = s.shapes.add_table(nr, nc, Inches(0.72), Inches(top), Inches(11.9),
                            Inches(min(5.2, 0.42 * nr))).table
    for j, wdt in enumerate(widths):
        gt.columns[j].width = Emu(int(Inches(11.9) * wdt / sum(widths)))
    for j, htxt in enumerate(header):
        c = gt.cell(0, j); c.text = ""
        p = c.text_frame.paragraphs[0]
        _run(p, htxt, hfs, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        c.fill.solid(); c.fill.fore_color.rgb = ACC
        c.margin_left = c.margin_right = Inches(0.1)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.text = ""
            p = c.text_frame.paragraphs[0]
            bold = (j == 0)
            _run(p, str(val), fs, bold=bold)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF4, 0xF6, 0xF8)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.04)
    _notes(s, notes)
    return s

def build(path, figs):
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    return prs
